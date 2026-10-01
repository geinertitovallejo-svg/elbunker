import streamlit as st
import pandas as pd
import os
from datetime import datetime
from io import BytesIO

# Configuración de la página
st.set_page_config(page_title="El Búnker - Sistema de Autolavado", layout="wide")
st.markdown(
    """
    <style>
    h1, h2, h3 {
        font-weight: 700;
    }
    [data-testid="stTextInput"] label,
    [data-testid="stSelectbox"] label,
    [data-testid="stNumberInput"] label {
        color: #aaa;
        font-weight: 500;
    }
    [data-testid="stVerticalBlockBorderWrapper"] {
        background: rgba(128, 139, 150, 0.07);
        border: 1px solid rgba(128, 139, 150, 0.24);
        border-radius: 12px;
    }
    .st-key-ingresar_admin button,
    .st-key-registrar_lavador_btn button,
    .st-key-anadir_tipo_button button,
    .st-key-anadir_servicio_button button,
    .st-key-guardar_registro_btn button {
        background-color: #F8D7DA !important;
        border-color: #E9B9BE !important;
        color: #842029 !important;
        -webkit-text-fill-color: #842029 !important;
    }
    .st-key-ingresar_admin button:hover,
    .st-key-registrar_lavador_btn button:hover,
    .st-key-anadir_tipo_button button:hover,
    .st-key-anadir_servicio_button button:hover,
    .st-key-guardar_registro_btn button:hover {
        background-color: #F1C2C7 !important;
        border-color: #DFA4AB !important;
    }
    .st-key-dar_baja_lavador_btn button,
    .st-key-quitar_tipo_button button,
    .st-key-quitar_servicio_button button,
    .st-key-cancelar_edicion_btn button,
    .st-key-eliminar_registro_btn button {
        background-color: #DDF2FF !important;
        border-color: #B9DDF2 !important;
        color: #075985 !important;
        -webkit-text-fill-color: #075985 !important;
    }
    .st-key-dar_baja_lavador_btn button:hover,
    .st-key-quitar_tipo_button button:hover,
    .st-key-quitar_servicio_button button:hover,
    .st-key-cancelar_edicion_btn button:hover,
    .st-key-eliminar_registro_btn button:hover {
        background-color: #C6E8FA !important;
        border-color: #A4D5EF !important;
    }
    [data-testid="stTextInput"] input,
    [data-testid="stNumberInput"] input,
    [data-testid="stDateInput"] input,
    [data-testid="stTimeInput"] input {
        background-color: #fff !important;
        color: #111 !important;
        -webkit-text-fill-color: #111;
    }
    [data-testid="stSelectbox"] [role="group"],
    [data-testid="stMultiSelect"] [role="group"] {
        background-color: #fff !important;
        color: #111 !important;
        border-radius: 6px;
        overflow: hidden;
    }
    [data-testid="stSelectbox"] [role="combobox"],
    [data-testid="stSelectbox"] input,
    [data-testid="stMultiSelect"] [role="combobox"],
    [data-testid="stMultiSelect"] input {
        background-color: #fff !important;
        color: #111 !important;
        -webkit-text-fill-color: #111;
    }
    [data-testid="stSelectbox"] [role="group"] button,
    [data-testid="stMultiSelect"] [role="group"] button {
        background-color: #B3262D !important;
        border-color: #B3262D !important;
        color: #fff !important;
    }
    [data-testid="stSelectbox"] [role="group"] button svg,
    [data-testid="stMultiSelect"] [role="group"] button svg {
        color: #fff !important;
        stroke: #fff !important;
    }
    [data-testid="stSelectbox"] [role="group"] button:hover,
    [data-testid="stMultiSelect"] [role="group"] button:hover {
        background-color: #92222A !important;
        border-color: #92222A !important;
    }
    [data-testid="stTextInput"] input::placeholder,
    [data-testid="stNumberInput"] input::placeholder,
    [data-testid="stDateInput"] input::placeholder,
    [data-testid="stTimeInput"] input::placeholder {
        color: #555 !important;
        -webkit-text-fill-color: #555;
    }
    [role="listbox"],
    [role="option"] {
        background-color: #fff !important;
        color: #111 !important;
    }
    [role="option"][aria-selected="true"] {
        background-color: #e9eef5 !important;
        color: #111 !important;
    }
    [role="option"]:hover {
        background-color: #B3262D !important;
        color: #fff !important;
    }
    [data-testid="stSelectboxVirtualDropdown"] {
        max-height: min(220px, 24vh) !important;
        overflow-y: auto !important;
    }
    .st-key-descargar_registros_excel button,
    [data-testid="stDownloadButton"] button {
        background-color: #198754 !important;
        border-color: #198754 !important;
        color: #fff !important;
    }
    .st-key-descargar_registros_excel button:hover,
    [data-testid="stDownloadButton"] button:hover {
        background-color: #146c43 !important;
        border-color: #146c43 !important;
    }
    body:has(#registros-theme-marker) [data-testid="stDataFrame"] canvas[style*="height: 36px"] {
        filter: sepia(1) saturate(6) hue-rotate(314deg) brightness(2.8);
    }
    </style>
    """,
    unsafe_allow_html=True,
)

ARCHIVO_CSV = "registros_lavado.csv"
ARCHIVO_LAVADORES = "lavadores.txt"
ARCHIVO_VEHICULOS = "vehiculos.txt"
ARCHIVO_SERVICIOS = "servicios.txt"
COLUMNAS_VALIDAS = ["ID", "Tipo", "Placa", "Servicio", "Pago", "Monto", "Fecha", "Lavador 1", "Lavador 2", "Motivo Gasto", "Precio Gasto"]
COLUMNAS_EXPORTACION = [
    ("Tipo", "TIPO"),
    ("Placa", "PLACA"),
    ("Servicio", "SERVICIO"),
    ("Pago", "PAGO"),
    ("Monto", "MONTO"),
    ("Fecha", "FECHA"),
    ("Lavador 1", "LAVADOR 1"),
    ("Lavador 2", "LAVADOR 2"),
    ("Motivo Gasto", "MOTIVO GASTO"),
    ("Precio Gasto", "PRECIO GASTO"),
]

def cargar_catalogo(archivo, valores_iniciales):
    if os.path.exists(archivo):
        with open(archivo, "r", encoding="utf-8") as catalogo:
            valores = [linea.strip() for linea in catalogo if linea.strip()]
        return valores
    return valores_iniciales.copy()

def guardar_catalogo(archivo, valores):
    with open(archivo, "w", encoding="utf-8") as catalogo:
        catalogo.write("\n".join(valores) + "\n")

def cargar_datos():
    if os.path.exists(ARCHIVO_CSV):
        try:
            df = pd.read_csv(ARCHIVO_CSV, dtype={'ID': str, 'Placa': str})
            if not all(col in df.columns for col in COLUMNAS_VALIDAS):
                return pd.DataFrame(columns=COLUMNAS_VALIDAS)
            return df
        except Exception:
            return pd.DataFrame(columns=COLUMNAS_VALIDAS)
    else:
        df_inicial = pd.DataFrame(columns=COLUMNAS_VALIDAS)
        df_inicial.to_csv(ARCHIVO_CSV, index=False)
        return df_inicial

def limpiar_monto(valor):
    try:
        val_str = str(valor).replace("S/", "").strip()
        if val_str in ["", "-", "nan", "None"]:
            return 0.0
        return float(val_str)
    except:
        return 0.0

def estilo_metodo_pago(valor):
    estilos = {
        "efectivo": "background-color: #DDF3E4; color: #245B35; font-weight: 600; text-align: center;",
        "tarjeta": "background-color: #DCEBFA; color: #1E4E79; font-weight: 600; text-align: center;",
        "yape/plin": "background-color: #EADFF5; color: #593477; font-weight: 600; text-align: center;",
    }
    return estilos.get(str(valor).strip().casefold())

def estilo_fila_tabla(fila, filas_seleccionadas):
    estilo_base = "background-color: #FFFFFF; color: #111111;"
    if fila.name in filas_seleccionadas:
        estilo_base = "background-color: rgba(179, 38, 45, 0.72); color: #FFFFFF;"
    estilos_celdas = [estilo_base] * len(fila)

    estilo_pago = estilo_metodo_pago(fila["Pago"])
    if estilo_pago:
        for columna in ("Pago", "Monto"):
            estilos_celdas[fila.index.get_loc(columna)] = estilo_pago

    return estilos_celdas

def crear_tabla_estilizada(df, filas_seleccionadas=None):
    df_visual = (
        df.drop(columns=["ID"], errors="ignore")
        .fillna("-")
        .astype(str)
        .reset_index(drop=True)
    )
    filas_seleccionadas = filas_seleccionadas or set()
    return df_visual.style.apply(
        lambda fila: estilo_fila_tabla(fila, filas_seleccionadas), axis=1
    )

def guardar_seleccion_tabla():
    estado_tabla = st.session_state.get("tabla_admin")
    seleccion = estado_tabla.get("selection", {}) if estado_tabla else {}
    st.session_state.filas_tabla_seleccionadas = set(seleccion.get("rows", []))

def parsear_fecha(valor):
    if pd.isna(valor):
        return pd.NaT

    texto = str(valor).strip()
    formatos = [
        "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y/%m/%d",
        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
        "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d-%m-%Y",
    ]
    for formato in formatos:
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    return pd.NaT

def calcular_resumen_financiero(df, ahora=None):
    if df.empty:
        return 0.0, 0.0, 0.0, 0.0

    df_metricas = df.copy()
    df_metricas["Fecha_Normalizada"] = df_metricas["Fecha"].map(parsear_fecha)
    df_metricas["Monto_Num"] = df_metricas["Monto"].map(limpiar_monto)
    df_metricas["Precio_Gasto_Num"] = df_metricas["Precio Gasto"].map(limpiar_monto).abs()

    tipo_gasto = df_metricas["Tipo"].fillna("").astype(str).str.contains(
        r"\b(?:gasto|egreso|expense)\b", case=False, regex=True
    )
    df_metricas["Ingreso_Num"] = df_metricas["Monto_Num"].where(~tipo_gasto, 0.0)
    df_metricas["Gasto_Num"] = df_metricas["Precio_Gasto_Num"].where(
        df_metricas["Precio_Gasto_Num"] > 0,
        df_metricas["Monto_Num"].abs().where(tipo_gasto, 0.0),
    )

    ahora = ahora or datetime.now()
    fechas = df_metricas["Fecha_Normalizada"]
    es_hoy = fechas.dt.date == ahora.date()
    es_mes_actual = (fechas.dt.year == ahora.year) & (fechas.dt.month == ahora.month)

    ganancia_dia = df_metricas.loc[es_hoy, "Ingreso_Num"].sum()
    gasto_dia = df_metricas.loc[es_hoy, "Gasto_Num"].sum()
    ganancia_mes = df_metricas.loc[es_mes_actual, "Ingreso_Num"].sum()
    gasto_mes = df_metricas.loc[es_mes_actual, "Gasto_Num"].sum()
    return ganancia_dia, gasto_dia, ganancia_mes, gasto_mes

def generar_excel_registros(df):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    df_exportacion = df[[columna for columna, _ in COLUMNAS_EXPORTACION]].copy()
    df_exportacion.columns = [titulo for _, titulo in COLUMNAS_EXPORTACION]

    for columna in ("MONTO", "PRECIO GASTO"):
        df_exportacion[columna] = df_exportacion[columna].map(
            lambda valor: "-"
            if str(valor).strip().casefold() in {"", "-", "nan", "none"}
            else limpiar_monto(valor)
        )

    archivo = BytesIO()
    with pd.ExcelWriter(archivo, engine="openpyxl") as writer:
        df_exportacion.to_excel(writer, index=False, sheet_name="Registros")
        hoja = writer.sheets["Registros"]
        formato_moneda = '"S/ " #,##0.00;"S/ " -#,##0.00;"S/ " 0.00'
        borde_fino = Side(style="thin", color="FFB7C9D6")
        borde_celda = Border(
            left=borde_fino,
            right=borde_fino,
            top=borde_fino,
            bottom=borde_fino,
        )
        relleno_cabecera = PatternFill(fill_type="solid", fgColor="FF000000")
        fuente_cabecera = Font(color="FFFFFFFF", bold=True)

        for celda in hoja[1]:
            celda.fill = relleno_cabecera
            celda.font = fuente_cabecera
            celda.alignment = Alignment(horizontal="center", vertical="center")

        for fila in hoja.iter_rows():
            for celda in fila:
                celda.border = borde_celda
                if celda.row > 1:
                    celda.alignment = Alignment(vertical="center")

        for numero_columna in (5, 10):
            for fila in range(2, hoja.max_row + 1):
                celda = hoja.cell(row=fila, column=numero_columna)
                if isinstance(celda.value, (int, float)):
                    celda.number_format = formato_moneda

        for celdas_columna in hoja.columns:
            letra_columna = celdas_columna[0].column_letter
            longitud_maxima = 0
            for celda in celdas_columna:
                if celda.value is None:
                    continue
                if celda.column in (5, 10) and isinstance(celda.value, (int, float)):
                    texto = f"S/ {celda.value:,.2f}"
                else:
                    texto = str(celda.value)
                longitud_maxima = max(longitud_maxima, len(texto))
            hoja.column_dimensions[letra_columna].width = max(longitud_maxima + 2, 12)

        hoja.row_dimensions[1].height = 24
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions

    return archivo.getvalue()

# --- CARGAR DATOS INICIALES ---
if "df_registros_memoria" not in st.session_state:
    st.session_state.df_registros_memoria = cargar_datos()
df_registros = st.session_state.df_registros_memoria

# Catálogos persistentes compartidos por el panel y el formulario
lista_lavadores = cargar_catalogo(
    ARCHIVO_LAVADORES, ["Carlos", "Junior", "Pedro", "Juan", "Jose"]
)
lista_tipos = cargar_catalogo(
    ARCHIVO_VEHICULOS, ["Auto", "Moto", "Motaxi", "Camioneta", "Bicicleta"]
)
lista_servicios = cargar_catalogo(
    ARCHIVO_SERVICIOS, ["Lavado Simple", "Lavado Completo", "Encerado", "Lavado De Salon"]
)
if not lista_tipos:
    lista_tipos = ["Auto"]
if not lista_servicios:
    lista_servicios = ["Lavado Simple"]
lista_pagos = ["Efectivo", "Tarjeta", "Yape/Plin"]
l1_opts = ["-", *lista_lavadores]
l2_opts = ["-", *lista_lavadores]

if "edit_id" not in st.session_state: st.session_state.edit_id = None
if "registro_cargado_id" not in st.session_state: st.session_state.registro_cargado_id = None
if "input_tipo" not in st.session_state: st.session_state.input_tipo = "Auto"
if "input_placa" not in st.session_state: st.session_state.input_placa = ""
if "input_servicio" not in st.session_state: st.session_state.input_servicio = "Lavado Simple"
if "input_monto" not in st.session_state: st.session_state.input_monto = "16.00"
if "input_pago" not in st.session_state: st.session_state.input_pago = "Efectivo"
if "input_l1" not in st.session_state: st.session_state.input_l1 = "Carlos"
if "input_l2" not in st.session_state: st.session_state.input_l2 = "-"
if "input_motivo" not in st.session_state: st.session_state.input_motivo = "-"
if "input_precio_gasto" not in st.session_state: st.session_state.input_precio_gasto = "-"

if st.session_state.input_tipo not in lista_tipos:
    st.session_state.input_tipo = lista_tipos[0] if lista_tipos else ""
if st.session_state.input_servicio not in lista_servicios:
    st.session_state.input_servicio = lista_servicios[0] if lista_servicios else ""
if st.session_state.input_l1 not in l1_opts:
    st.session_state.input_l1 = "Carlos" if "Carlos" in lista_lavadores else (lista_lavadores[0] if lista_lavadores else "-")
if st.session_state.input_l2 not in l2_opts:
    st.session_state.input_l2 = "-"

if st.session_state.pop("resetear_formulario", False):
    st.session_state.input_tipo = lista_tipos[0] if lista_tipos else ""
    st.session_state.input_placa = ""
    st.session_state.input_servicio = lista_servicios[0] if lista_servicios else ""
    st.session_state.input_monto = "16.00"
    st.session_state.input_pago = "Efectivo"
    st.session_state.input_l1 = "Carlos" if "Carlos" in lista_lavadores else (lista_lavadores[0] if lista_lavadores else "-")
    st.session_state.input_l2 = "-"
    st.session_state.input_motivo = "-"
    st.session_state.input_precio_gasto = "-"

def limpiar_formulario():
    st.session_state.edit_id = None
    st.session_state.resetear_formulario = True

def cargar_fila_en_formulario(fila_data):
    registro_id = str(fila_data['ID'])
    if st.session_state.registro_cargado_id == registro_id:
        return

    st.session_state.registro_cargado_id = registro_id
    st.session_state.edit_id = registro_id
    st.session_state.input_tipo = fila_data['Tipo'] if fila_data['Tipo'] in lista_tipos else lista_tipos[0]
    st.session_state.input_placa = str(fila_data['Placa'])
    st.session_state.input_servicio = fila_data['Servicio'] if fila_data['Servicio'] in lista_servicios else lista_servicios[0]
    st.session_state.input_monto = str(fila_data['Monto']).replace("S/", "").strip()
    st.session_state.input_pago = fila_data['Pago'] if fila_data['Pago'] in lista_pagos else "Efectivo"
    st.session_state.input_l1 = fila_data['Lavador 1'] if fila_data['Lavador 1'] in l1_opts else (
        "Carlos" if "Carlos" in lista_lavadores else (lista_lavadores[0] if lista_lavadores else "-")
    )
    st.session_state.input_l2 = fila_data['Lavador 2'] if fila_data['Lavador 2'] in l2_opts else "-"
    st.session_state.input_motivo = str(fila_data['Motivo Gasto'])
    st.session_state.input_precio_gasto = str(fila_data['Precio Gasto']).replace("S/", "").strip()

# --- NAVEGACIÓN Y VISTAS ---
opciones_navegacion = ["Registros", "Caja", "Personal", "Ajustes"]
if st.session_state.get("seccion_navegacion") not in opciones_navegacion:
    st.session_state.seccion_navegacion = "Registros"

st.sidebar.markdown("## The Bunker")
st.sidebar.markdown(
    """
    <style>
    [data-testid="stSidebar"] [data-testid^="stBaseButton-"] {
        min-height: 46px;
        border-radius: 12px;
        justify-content: flex-start;
        padding-left: 1rem;
        font-weight: 600;
        transition: background-color 140ms ease, border-color 140ms ease;
    }
    [data-testid="stSidebar"] [data-testid="stBaseButton-secondary"] {
        border: 1px solid rgba(128, 128, 128, 0.35);
    }
    </style>
    """,
    unsafe_allow_html=True,
)

def seleccionar_seccion(nombre):
    st.session_state.seccion_navegacion = nombre

for opcion in opciones_navegacion:
    st.sidebar.button(
        opcion,
        key=f"navegacion_{opcion.lower()}",
        type="primary" if st.session_state.seccion_navegacion == opcion else "secondary",
        use_container_width=True,
        on_click=seleccionar_seccion,
        args=(opcion,),
    )

seccion = st.session_state.seccion_navegacion
if "admin_autorizado" not in st.session_state:
    st.session_state.admin_autorizado = False

is_admin = st.session_state.admin_autorizado
if seccion != "Registros" and not is_admin:
    st.title(seccion)
    st.subheader("Acceso de administrador")
    password_admin = st.text_input("Contraseña Admin", type="password", key="password_admin")
    if st.button("Ingresar", key="ingresar_admin", type="primary"):
        if password_admin == "1234":
            st.session_state.admin_autorizado = True
            st.rerun()
        else:
            st.error("Contraseña incorrecta.")
    st.stop()

if seccion == "Caja":
    st.title("Caja")
    ganancia_dia, gasto_dia, ganancia_mes, gasto_mes = calcular_resumen_financiero(
        df_registros, datetime.now()
    )
    st.subheader("Balance de Ingresos")
    col_ingreso_dia, col_ingreso_mes = st.columns(2)
    col_ingreso_dia.metric("Ingresos del día", f"S/ {ganancia_dia:.2f}")
    col_ingreso_mes.metric("Ingresos del mes actual", f"S/ {ganancia_mes:.2f}")
    st.subheader("Balance de Gastos")
    col_gasto_dia, col_gasto_mes = st.columns(2)
    col_gasto_dia.metric("Gastos del día", f"S/ {gasto_dia:.2f}")
    col_gasto_mes.metric("Gastos del mes actual", f"S/ {gasto_mes:.2f}")
    st.stop()

if seccion == "Personal":
    st.title("Personal")
    st.subheader("Control de Lavadores")
    if lista_lavadores:
        lavador_seleccionado = st.selectbox(
            "Personal registrado", lista_lavadores, key="admin_lavador_seleccionado"
        )
    else:
        lavador_seleccionado = None
        st.info("No hay lavadores registrados.")

    nuevo_lavador = st.text_input("Nuevo nombre", key="nuevo_lavador_input")
    col_registrar_lavador, col_baja_lavador = st.columns(2)
    if col_registrar_lavador.button(
        "Registrar Lavador", key="registrar_lavador_btn", type="primary", use_container_width=True
    ):
        nombre = nuevo_lavador.strip()
        if not nombre:
            st.warning("Escribe el nombre del lavador.")
        elif any(nombre.casefold() == existente.casefold() for existente in lista_lavadores):
            st.warning("Ese lavador ya está registrado.")
        else:
            guardar_catalogo(ARCHIVO_LAVADORES, [*lista_lavadores, nombre])
            st.success(f"Lavador {nombre} registrado.")
            st.rerun()

    if col_baja_lavador.button(
        "Dar de Baja Seleccionado", key="dar_baja_lavador_btn", type="secondary",
        use_container_width=True, disabled=lavador_seleccionado is None
    ):
        lista_actualizada = [
            nombre for nombre in lista_lavadores if nombre != lavador_seleccionado
        ]
        guardar_catalogo(ARCHIVO_LAVADORES, lista_actualizada)
        for clave in ("input_l1", "input_l2"):
            if st.session_state.get(clave) == lavador_seleccionado:
                st.session_state[clave] = "-"
        st.success(f"Lavador {lavador_seleccionado} dado de baja.")
        st.rerun()
    st.stop()

if seccion == "Ajustes":
    st.title("Ajustes")
    st.subheader("Tipos de Vehículos")
    tipo_seleccionado = st.selectbox(
        "Tipos registrados", lista_tipos, key="admin_tipo_seleccionado"
    )
    nuevo_tipo = st.text_input("Nuevo tipo de vehículo", key="nuevo_tipo_input")
    col_anadir_tipo, col_quitar_tipo = st.columns(2)
    if col_anadir_tipo.button(
        "Añadir", key="anadir_tipo_button", type="primary", use_container_width=True
    ):
        tipo = nuevo_tipo.strip()
        if not tipo:
            st.warning("Escribe el tipo de vehículo que deseas añadir.")
        elif any(tipo.casefold() == existente.casefold() for existente in lista_tipos):
            st.warning("Ese tipo de vehículo ya está registrado.")
        else:
            guardar_catalogo(ARCHIVO_VEHICULOS, [*lista_tipos, tipo])
            st.success(f"Tipo de vehículo {tipo} añadido.")
            st.rerun()
    if col_quitar_tipo.button(
        "Quitar", key="quitar_tipo_button", type="secondary", use_container_width=True,
        disabled=len(lista_tipos) <= 1
    ):
        tipos_actualizados = [tipo for tipo in lista_tipos if tipo != tipo_seleccionado]
        guardar_catalogo(ARCHIVO_VEHICULOS, tipos_actualizados)
        if st.session_state.input_tipo == tipo_seleccionado:
            st.session_state.input_tipo = tipos_actualizados[0]
        st.success(f"Tipo de vehículo {tipo_seleccionado} quitado.")
        st.rerun()

    st.markdown("---")
    st.subheader("Tipos de Servicios")
    servicio_seleccionado = st.selectbox(
        "Servicios registrados", lista_servicios, key="admin_servicio_seleccionado"
    )
    nuevo_servicio = st.text_input("Nuevo servicio", key="nuevo_servicio_input")
    col_anadir_servicio, col_quitar_servicio = st.columns(2)
    if col_anadir_servicio.button(
        "Añadir", key="anadir_servicio_button", type="primary", use_container_width=True
    ):
        servicio = nuevo_servicio.strip()
        if not servicio:
            st.warning("Escribe el servicio que deseas añadir.")
        elif any(servicio.casefold() == existente.casefold() for existente in lista_servicios):
            st.warning("Ese servicio ya está registrado.")
        else:
            guardar_catalogo(ARCHIVO_SERVICIOS, [*lista_servicios, servicio])
            st.success(f"Servicio {servicio} añadido.")
            st.rerun()
    if col_quitar_servicio.button(
        "Quitar", key="quitar_servicio_button", type="secondary", use_container_width=True,
        disabled=len(lista_servicios) <= 1
    ):
        servicios_actualizados = [
            servicio for servicio in lista_servicios if servicio != servicio_seleccionado
        ]
        guardar_catalogo(ARCHIVO_SERVICIOS, servicios_actualizados)
        if st.session_state.input_servicio == servicio_seleccionado:
            st.session_state.input_servicio = servicios_actualizados[0]
        st.success(f"Servicio {servicio_seleccionado} quitado.")
        st.rerun()
    st.stop()

st.markdown('<div id="registros-theme-marker"></div>', unsafe_allow_html=True)
st.title("Registros")

# --- TABLA DE REGISTROS (PROCESADA PRIMERO PARA CARGAR LA SELECCIÓN) ---
df_exportacion = df_registros.copy()
if is_admin and not df_registros.empty:
    st.subheader("Registros del Día (Selecciona una fila para editar)")

    with st.container(border=True):
        col_f1, col_f2 = st.columns(2)
        with col_f1:
            busqueda_placa = st.text_input("Buscar por Placa", key="busqueda_placa_input")
        with col_f2:
            filtro_pago = st.selectbox("Filtrar por Método de Pago", ["Todos", "Efectivo", "Tarjeta", "Yape/Plin"], key="filtro_pago_select")
        
    df_filtrado = df_registros.copy()
    if busqueda_placa:
        df_filtrado = df_filtrado[df_filtrado['Placa'].astype(str).str.contains(busqueda_placa, case=False, na=False)]
    if filtro_pago != "Todos":
        df_filtrado = df_filtrado[df_filtrado['Pago'] == filtro_pago]
    df_exportacion = df_filtrado.copy()

    filas_seleccionadas = st.session_state.get("filas_tabla_seleccionadas", set())

    with st.container(border=True):
        evento_tabla = st.dataframe(
            crear_tabla_estilizada(df_filtrado, filas_seleccionadas),
            use_container_width=True,
            selection_mode="single-row",
            on_select=guardar_seleccion_tabla,
            key="tabla_admin",
            hide_index=True,
        )
    
    # Capturar la selección de la tabla y actualizar session_state ANTES de crear los inputs
    if evento_tabla and "selection" in evento_tabla and evento_tabla["selection"].get("rows"):
        indice_seleccionado = evento_tabla["selection"]["rows"][0]
        if len(df_filtrado) > indice_seleccionado:
            fila_data = df_filtrado.iloc[indice_seleccionado]
            cargar_fila_en_formulario(fila_data)
    elif not st.session_state.edit_id:
        st.session_state.registro_cargado_id = None

if not is_admin and not df_registros.empty:
    st.subheader("Registros del Día")
    with st.container(border=True):
        st.dataframe(
            crear_tabla_estilizada(df_registros),
            use_container_width=True,
            hide_index=True,
        )
elif df_registros.empty:
    st.info("Aún no hay registros en el archivo CSV.")

st.download_button(
    "Exportar a Excel",
    data=generar_excel_registros(df_exportacion),
    file_name=f"registros_{datetime.now().strftime('%Y-%m-%d')}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    key="descargar_registros_excel",
    type="primary",
    use_container_width=True,
)

st.markdown("---")

# --- FORMULARIO DE REGISTRO / ACTUALIZACIÓN ---
titulo_form = "Actualizar Registro Seleccionado" if st.session_state.edit_id else "Registrar Nuevo Servicio / Gasto"
st.subheader(titulo_form)

if st.session_state.edit_id:
    col_info1, col_info2 = st.columns([4, 1])
    with col_info1:
        st.warning(f"Modificando el registro de la placa: **{st.session_state.input_placa}**. Modifica los campos abajo y haz clic en Guardar Cambios, o cancela.")
    with col_info2:
        if st.button(
            "Cancelar", key="cancelar_edicion_btn", type="secondary", use_container_width=True
        ):
            limpiar_formulario()
            st.rerun()

with st.container(border=True):
    st.markdown("### Vehículo y Pago")
    c1, c2, c3, c4, c5 = st.columns(5)
    
    with c1:
        st.selectbox("Tipo de Vehículo", lista_tipos, key="input_tipo")
    with c2:
        st.text_input("Placa / Identificador", placeholder="Ej. ABC-123", key="input_placa")
    with c3:
        st.selectbox("Servicio", lista_servicios, key="input_servicio")
    with c4:
        st.text_input("Monto Cobrado (S/)", key="input_monto")
    with c5:
        st.selectbox("Forma de Pago", lista_pagos, key="input_pago")
        
with st.container(border=True):
    st.markdown("### Lavadores")
    l1, l2 = st.columns(2)
    with l1:
        st.selectbox("Lavador 1", l1_opts, key="input_l1")
    with l2:
        st.selectbox("Lavador 2", l2_opts, key="input_l2")
        
with st.container(border=True):
    st.markdown("### Gastos")
    g1, g2 = st.columns(2)
    with g1:
        st.text_input("Motivo Gasto (opcional)", key="input_motivo")
    with g2:
        st.text_input("Precio Gasto (opcional)", key="input_precio_gasto")

col_acc1, col_acc2 = st.columns(2)
with col_acc1:
    texto_boton = "Guardar Cambios de Actualización" if st.session_state.edit_id else "Guardar Registro"
    btn_guardar = st.button(
        texto_boton, key="guardar_registro_btn", use_container_width=True, type="primary"
    )
with col_acc2:
    if st.session_state.edit_id:
        btn_eliminar = st.button(
            "Eliminar este Registro", key="eliminar_registro_btn", type="secondary",
            use_container_width=True
        )
    else:
        btn_eliminar = False

# Lógica de Guardar / Actualizar
if btn_guardar:
    if st.session_state.input_placa.strip() != "":
        df_actual = df_registros.copy()
        m_val = st.session_state.input_monto
        pg_val = st.session_state.input_precio_gasto
        
        monto_fmt = f"S/ {float(m_val):.2f}" if str(m_val).replace('.','',1).isdigit() else f"S/ {m_val}"
        gasto_fmt = f"S/ {float(pg_val):.2f}" if str(pg_val).replace('.','',1).isdigit() else pg_val
        
        if st.session_state.edit_id:
            mask = df_actual['ID'].astype(str) == str(st.session_state.edit_id)
            if mask.any():
                df_actual.loc[mask, 'Tipo'] = st.session_state.input_tipo
                df_actual.loc[mask, 'Placa'] = st.session_state.input_placa.upper()
                df_actual.loc[mask, 'Servicio'] = st.session_state.input_servicio
                df_actual.loc[mask, 'Pago'] = st.session_state.input_pago
                df_actual.loc[mask, 'Monto'] = monto_fmt
                df_actual.loc[mask, 'Lavador 1'] = st.session_state.input_l1
                df_actual.loc[mask, 'Lavador 2'] = st.session_state.input_l2
                df_actual.loc[mask, 'Motivo Gasto'] = st.session_state.input_motivo
                df_actual.loc[mask, 'Precio Gasto'] = gasto_fmt
                
                df_actual.to_csv(ARCHIVO_CSV, index=False)
                st.session_state.df_registros_memoria = df_actual
                limpiar_formulario()
                st.success("¡Registro actualizado en tiempo real con éxito!")
                st.rerun()
            else:
                st.error("No se encontró el registro a actualizar.")
        else:
            id_unico = str(int(datetime.now().timestamp() * 1000))
            fecha_actual = datetime.now().strftime("%Y/%m/%d %H:%M")
            
            nuevo_registro = pd.DataFrame({
                "ID": [id_unico],
                "Tipo": [st.session_state.input_tipo],
                "Placa": [st.session_state.input_placa.upper()],
                "Servicio": [st.session_state.input_servicio],
                "Pago": [st.session_state.input_pago],
                "Monto": [monto_fmt],
                "Fecha": [fecha_actual],
                "Lavador 1": [st.session_state.input_l1],
                "Lavador 2": [st.session_state.input_l2],
                "Motivo Gasto": [st.session_state.input_motivo],
                "Precio Gasto": [gasto_fmt]
            })
            
            df_actual = pd.concat([df_actual, nuevo_registro], ignore_index=True)
            df_actual.to_csv(ARCHIVO_CSV, index=False)
            st.session_state.df_registros_memoria = df_actual
            limpiar_formulario()
            st.success("¡Registro guardado con éxito!")
            st.rerun()
    else:
        st.warning("Por favor, ingresa al menos la placa o identificador.")

if btn_eliminar and st.session_state.edit_id:
    df_nuevo = df_registros[df_registros['ID'].astype(str) != str(st.session_state.edit_id)].copy()
    df_nuevo.to_csv(ARCHIVO_CSV, index=False)
    st.session_state.df_registros_memoria = df_nuevo
    limpiar_formulario()
    st.success("¡Registro eliminado en tiempo real correctamente!")
    st.rerun()
